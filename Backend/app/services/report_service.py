from datetime import datetime, timezone
import io
from typing import Any, Dict, List, Optional
import uuid
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.models.customer import Customer
from app.models.enums import (
    CustomerStage,
    KedlDemandStatus,
    LeadStatus,
    PaymentMode,
    PaymentStatus,
    TicketStatus,
    UserRole,
    WorkStatus,
)
from app.models.inventory import InventoryItem, SerialItem, StockMovement
from app.models.kedl import KedlDemand, KedlFile
from app.models.lead import Lead
from app.models.payment import Payment
from app.models.team import Team
from app.models.ticket import ServiceTicket
from app.models.user import User
from app.models.work_assignment import WorkAssignment
from app.schemas.report import (
    DashboardReportResponse,
    KedlReportResponse,
    PaymentsReportResponse,
    SalesReportResponse,
    SalesRepPerformance,
    TeamWorkMetrics,
    WorkReportResponse,
)


class ReportService:
    @staticmethod
    async def get_dashboard_report(db: AsyncSession) -> DashboardReportResponse:
        # 1. Leads by status
        lead_stmt = select(Lead.status, func.count(Lead.id)).where(Lead.is_deleted == False).group_by(Lead.status)  # noqa: E712
        lead_rows = (await db.execute(lead_stmt)).all()
        leads_by_status = {r[0].value if hasattr(r[0], "value") else str(r[0]): r[1] for r in lead_rows}
        total_leads = sum(leads_by_status.values())

        # 2. Customers by stage
        cust_stmt = select(Customer.stage, func.count(Customer.id)).where(Customer.is_deleted == False).group_by(Customer.stage)  # noqa: E712
        cust_rows = (await db.execute(cust_stmt)).all()
        customers_by_stage = {r[0].value if hasattr(r[0], "value") else str(r[0]): r[1] for r in cust_rows}
        total_customers = sum(customers_by_stage.values())

        # 3. Sales this month (paise)
        now = datetime.now(timezone.utc)
        month_start = datetime(now.year, now.month, 1, tzinfo=timezone.utc)
        sales_stmt = select(func.sum(Customer.final_price)).where(
            Customer.is_deleted == False,  # noqa: E712
            Customer.created_at >= month_start,
        )
        sales_this_month = (await db.execute(sales_stmt)).scalar_one() or 0

        # 4. Pending payment verifications
        pend_pay_stmt = select(func.count(Payment.id)).where(
            Payment.is_deleted == False,  # noqa: E712
            Payment.status == PaymentStatus.SALES_APPROVED,
        )
        pending_payments = (await db.execute(pend_pay_stmt)).scalar_one() or 0

        # 5. Active works count
        work_stmt = select(func.count(WorkAssignment.id)).where(
            WorkAssignment.status.in_([WorkStatus.PENDING, WorkStatus.IN_PROGRESS])
        )
        active_works = (await db.execute(work_stmt)).scalar_one() or 0

        # 6. Open KEDL demands count
        kedl_stmt = select(func.count(KedlDemand.id)).where(
            KedlDemand.status == KedlDemandStatus.OPEN
        )
        open_kedl_demands = (await db.execute(kedl_stmt)).scalar_one() or 0

        # 7. Open service tickets count
        ticket_stmt = select(func.count(ServiceTicket.id)).where(
            ServiceTicket.status.in_([
                TicketStatus.OPEN,
                TicketStatus.ASSIGNED,
                TicketStatus.IN_PROGRESS,
                TicketStatus.REOPENED,
            ])
        )
        open_tickets = (await db.execute(ticket_stmt)).scalar_one() or 0

        # 8. Low stock items count
        stock_stmt = select(func.count(InventoryItem.id)).where(
            InventoryItem.is_deleted == False,  # noqa: E712
            InventoryItem.is_active == True,  # noqa: E712
            InventoryItem.stock_qty <= InventoryItem.min_stock,
        )
        low_stock_count = (await db.execute(stock_stmt)).scalar_one() or 0

        return DashboardReportResponse(
            total_leads=total_leads,
            leads_by_status=leads_by_status,
            total_customers=total_customers,
            customers_by_stage=customers_by_stage,
            sales_this_month_paise=int(sales_this_month),
            pending_payment_verifications_count=pending_payments,
            active_works_count=active_works,
            open_kedl_demands_count=open_kedl_demands,
            open_service_tickets_count=open_tickets,
            low_stock_items_count=low_stock_count,
        )

    @staticmethod
    async def get_sales_report(db: AsyncSession) -> SalesReportResponse:
        # Get all sales reps
        reps_stmt = select(User).where(User.role == UserRole.SALES_EXECUTIVE, User.is_deleted == False)  # noqa: E712
        reps = list((await db.execute(reps_stmt)).scalars().all())

        sales_reps_metrics: List[SalesRepPerformance] = []
        grand_total_sales = 0

        for rep in reps:
            # Leads count
            lead_stmt = select(func.count(Lead.id)).where(Lead.assigned_to == rep.id, Lead.is_deleted == False)  # noqa: E712
            lead_count = (await db.execute(lead_stmt)).scalar_one() or 0

            # Converted customers & total sales amount
            cust_stmt = select(func.count(Customer.id), func.sum(Customer.final_price)).where(
                Customer.sales_rep_id == rep.id,
                Customer.is_deleted == False,  # noqa: E712
            )
            cust_row = (await db.execute(cust_stmt)).first()
            cust_count = cust_row[0] if cust_row else 0
            rep_sales = int(cust_row[1] or 0) if cust_row else 0

            grand_total_sales += rep_sales
            sales_reps_metrics.append(
                SalesRepPerformance(
                    sales_rep_id=rep.id,
                    sales_rep_name=rep.name,
                    total_leads=lead_count,
                    converted_customers=cust_count,
                    total_sales_paise=rep_sales,
                )
            )

        return SalesReportResponse(
            sales_reps=sales_reps_metrics,
            total_sales_paise=grand_total_sales,
        )

    @staticmethod
    async def get_payments_report(db: AsyncSession) -> PaymentsReportResponse:
        # Total billed
        billed_stmt = select(func.sum(Customer.final_price)).where(Customer.is_deleted == False)  # noqa: E712
        total_billed = (await db.execute(billed_stmt)).scalar_one() or 0

        # Total collected
        coll_stmt = select(func.sum(Payment.amount)).where(
            Payment.is_deleted == False,  # noqa: E712
            Payment.status == PaymentStatus.VERIFIED,
        )
        total_collected = (await db.execute(coll_stmt)).scalar_one() or 0

        # Payments by mode
        mode_stmt = select(Payment.mode, func.sum(Payment.amount)).where(
            Payment.is_deleted == False,  # noqa: E712
            Payment.status == PaymentStatus.VERIFIED,
        ).group_by(Payment.mode)
        mode_rows = (await db.execute(mode_stmt)).all()
        by_mode = {r[0].value if hasattr(r[0], "value") else str(r[0]): int(r[1] or 0) for r in mode_rows}

        pending_advance = max(0, int(total_billed * 0.2) - int(total_collected))
        pending_final = max(0, int(total_billed) - int(total_collected))

        return PaymentsReportResponse(
            total_billed_paise=int(total_billed),
            total_collected_paise=int(total_collected),
            pending_advance_paise=pending_advance,
            pending_final_paise=pending_final,
            payments_by_mode=by_mode,
        )

    @staticmethod
    async def get_work_report(db: AsyncSession) -> WorkReportResponse:
        teams_stmt = select(Team).where(Team.is_deleted == False)  # noqa: E712
        teams = list((await db.execute(teams_stmt)).scalars().all())

        team_metrics: List[TeamWorkMetrics] = []
        total_completed = 0
        total_active = 0

        for t in teams:
            status_stmt = select(WorkAssignment.status, func.count(WorkAssignment.id)).where(
                WorkAssignment.team_id == t.id
            ).group_by(WorkAssignment.status)
            rows = (await db.execute(status_stmt)).all()
            status_counts = {r[0]: r[1] for r in rows}

            comp = status_counts.get(WorkStatus.COMPLETED, 0)
            in_prog = status_counts.get(WorkStatus.IN_PROGRESS, 0)
            pend = status_counts.get(WorkStatus.PENDING, 0)

            total_completed += comp
            total_active += (in_prog + pend)

            team_metrics.append(
                TeamWorkMetrics(
                    team_id=t.id,
                    team_name=t.name,
                    team_type=t.type.value,
                    completed_works=comp,
                    in_progress_works=in_prog,
                    pending_works=pend,
                )
            )

        return WorkReportResponse(
            teams=team_metrics,
            total_completed=total_completed,
            total_active=total_active,
        )

    @staticmethod
    async def get_kedl_report(db: AsyncSession) -> KedlReportResponse:
        file_stmt = select(KedlFile.status, func.count(KedlFile.id)).group_by(KedlFile.status)
        f_rows = (await db.execute(file_stmt)).all()
        files_by_status = {r[0].value if hasattr(r[0], "value") else str(r[0]): r[1] for r in f_rows}

        open_dem_stmt = select(func.count(KedlDemand.id), func.sum(KedlDemand.amount)).where(
            KedlDemand.status == KedlDemandStatus.OPEN
        )
        open_row = (await db.execute(open_dem_stmt)).first()
        open_count = open_row[0] if open_row else 0
        open_amount = int(open_row[1] or 0) if open_row else 0

        paid_dem_stmt = select(func.sum(KedlDemand.amount)).where(
            KedlDemand.status == KedlDemandStatus.PAID
        )
        paid_amount = int((await db.execute(paid_dem_stmt)).scalar_one() or 0)

        return KedlReportResponse(
            files_by_status=files_by_status,
            open_demands_count=open_count,
            open_demands_total_paise=open_amount,
            paid_demands_total_paise=paid_amount,
        )

    # ------------------ EXCEL WORKBOOK EXPORTS ------------------
    @staticmethod
    async def export_customers_excel(db: AsyncSession) -> io.BytesIO:
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill
        except ImportError:
            raise RuntimeError("openpyxl is required for Excel export.")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Customers"

        # Headers
        headers = [
            "Customer ID",
            "Name",
            "Phone",
            "Address",
            "Capacity (kW)",
            "System Phase",
            "Panel Brand",
            "Inverter Brand",
            "Current Stage",
            "Final Price (INR)",
            "Sales Rep",
            "Created Date",
        ]
        ws.append(headers)

        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font

        # Query customers
        stmt = (
            select(Customer)
            .options(selectinload(Customer.sales_agent))
            .where(Customer.is_deleted == False)  # noqa: E712
            .order_by(Customer.created_at.desc())
        )
        customers = list((await db.execute(stmt)).scalars().all())

        for c in customers:
            ws.append([
                str(c.id),
                c.name,
                c.phone,
                c.address,
                float(c.capacity_kw),
                c.phase.value if c.phase else "",
                c.panel_brand or "",
                c.inverter_brand or "",
                c.stage.value if c.stage else "",
                float(c.final_price) / 100.0,
                c.sales_agent.name if c.sales_agent else "",
                c.created_at.strftime("%Y-%m-%d %H:%M"),
            ])

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @staticmethod
    async def export_payments_excel(db: AsyncSession) -> io.BytesIO:
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill
        except ImportError:
            raise RuntimeError("openpyxl is required for Excel export.")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Payments Ledger"

        headers = [
            "Payment ID",
            "Customer Name",
            "Customer Phone",
            "Amount (INR)",
            "Payment Mode",
            "Status",
            "Reference No",
            "Paid Date",
            "Recorded By",
        ]
        ws.append(headers)

        header_fill = PatternFill(start_color="047857", end_color="047857", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font

        stmt = (
            select(Payment)
            .options(selectinload(Payment.customer), selectinload(Payment.creator))
            .where(Payment.is_deleted == False)  # noqa: E712
            .order_by(Payment.created_at.desc())
        )
        payments = list((await db.execute(stmt)).scalars().all())

        for p in payments:
            ws.append([
                str(p.id),
                p.customer.name if p.customer else "",
                p.customer.phone if p.customer else "",
                float(p.amount) / 100.0,
                p.mode.value,
                p.status.value,
                p.reference_no or "",
                p.paid_on.strftime("%Y-%m-%d %H:%M") if p.paid_on else "",
                p.creator.name if p.creator else "",
            ])

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @staticmethod
    async def export_inventory_excel(db: AsyncSession) -> io.BytesIO:
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill
        except ImportError:
            raise RuntimeError("openpyxl is required for Excel export.")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Inventory Items"

        headers = [
            "Item ID",
            "Category",
            "Name",
            "Brand",
            "Model",
            "Unit",
            "Has Serials",
            "Min Safety Stock",
            "Current Stock Qty",
            "Stock Status",
        ]
        ws.append(headers)

        header_fill = PatternFill(start_color="B45309", end_color="B45309", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font

        stmt = (
            select(InventoryItem)
            .where(InventoryItem.is_deleted == False)  # noqa: E712
            .order_by(InventoryItem.category.asc(), InventoryItem.name.asc())
        )
        items = list((await db.execute(stmt)).scalars().all())

        for i in items:
            stock_status = "LOW STOCK" if i.stock_qty <= i.min_stock else "IN STOCK"
            ws.append([
                str(i.id),
                i.category.value,
                i.name,
                i.brand or "",
                i.model or "",
                i.unit.value,
                "YES" if i.has_serial else "NO",
                i.min_stock,
                i.stock_qty,
                stock_status,
            ])

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output
